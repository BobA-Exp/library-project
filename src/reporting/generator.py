"""Export catalog analytics to Excel, CSV, and JSON."""

import csv
import json
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import Settings, get_settings
from db.manager import DatabaseManager
from processing.analytics import DataProcessor

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_METRIC_COLUMNS = ("metric", "value")
_SUMMARY_COLUMNS = ("summary", "label", "work_count")


class ReportGenerator:
    """Write the formatted workbook, CSV summary, and minified JSON dump."""

    def __init__(
        self,
        processor: DataProcessor,
        settings: Settings | None = None,
    ) -> None:
        configured = settings or get_settings()
        self._processor = processor
        self._output_dir = Path(configured.report_output_dir)

    @classmethod
    def from_database(
        cls,
        database: DatabaseManager,
        settings: Settings | None = None,
    ) -> "ReportGenerator":
        processor = DataProcessor(database.list_works(), database.list_authors())
        return cls(processor, settings=settings)

    def generate(self, on: date | None = None) -> dict[str, Path]:
        """Write the three report files and return their paths."""
        stamp = (on or datetime.now(UTC).date()).strftime("%Y%m%d")
        self._output_dir.mkdir(parents=True, exist_ok=True)
        paths = {
            "workbook": self._output_dir / f"library_report_{stamp}.xlsx",
            "summary": self._output_dir / f"library_report_{stamp}.csv",
            "json": self._output_dir / f"library_report_{stamp}.json",
        }
        metrics = _records(self._processor.performance_metrics())
        years = _records(self._processor.top_publication_years())
        languages = _records(self._processor.top_languages())
        _write_workbook(paths["workbook"], metrics, years, languages)
        _write_summary(paths["summary"], years, languages)
        _write_json(paths["json"], metrics, years, languages)
        return paths


def _write_workbook(
    path: Path,
    metrics: list[dict[str, Any]],
    years: list[dict[str, Any]],
    languages: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    metrics_sheet = workbook.active
    metrics_sheet.title = "Performance Metrics"
    _write_sheet(metrics_sheet, metrics, list(_METRIC_COLUMNS))
    _write_sheet(
        workbook.create_sheet("Publication Years"),
        years,
        ["publication_year", "work_count"],
    )
    _write_sheet(
        workbook.create_sheet("Languages"),
        languages,
        ["language", "work_count"],
    )
    workbook.save(path)


def _write_sheet(
    sheet: Worksheet, rows: list[dict[str, Any]], columns: list[str]
) -> None:
    for column_index, name in enumerate(columns, start=1):
        cell = sheet.cell(1, column_index, name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for row_index, record in enumerate(rows, start=2):
        for column_index, name in enumerate(columns, start=1):
            sheet.cell(row_index, column_index, record.get(name))
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions
    for column_index, name in enumerate(columns, start=1):
        lengths = [len(name)]
        lengths.extend(len(str(record.get(name) or "")) for record in rows)
        width = min(40, max(12, max(lengths) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def _write_summary(
    path: Path,
    years: list[dict[str, Any]],
    languages: list[dict[str, Any]],
) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=_SUMMARY_COLUMNS)
        writer.writeheader()
        for record in years:
            writer.writerow(
                {
                    "summary": "publication_year",
                    "label": record.get("publication_year"),
                    "work_count": record.get("work_count"),
                }
            )
        for record in languages:
            writer.writerow(
                {
                    "summary": "language",
                    "label": record.get("language"),
                    "work_count": record.get("work_count"),
                }
            )


def _write_json(
    path: Path,
    metrics: list[dict[str, Any]],
    years: list[dict[str, Any]],
    languages: list[dict[str, Any]],
) -> None:
    payload = {
        "metrics": metrics,
        "publication_years": years,
        "languages": languages,
    }
    path.write_text(
        json.dumps(payload, separators=(",", ":"), ensure_ascii=False),
        encoding="utf-8",
    )


def _records(frame: pd.DataFrame) -> list[dict[str, Any]]:
    records = frame.to_dict(orient="records")
    return [{key: _plain(value) for key, value in row.items()} for row in records]


def _plain(value: Any) -> Any:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass
    item = getattr(value, "item", None)
    if callable(item):
        return item()
    return value
