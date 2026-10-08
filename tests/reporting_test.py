import csv
import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from config import Settings
from db.manager import DatabaseManager
from processing.analytics import DataProcessor
from reporting.generator import ReportGenerator

REPORT_DATE = date(2026, 10, 8)


def make_settings(tmp_path: Path) -> Settings:
    return Settings(
        _env_file=None,
        database_path=tmp_path / "library.db",
        report_output_dir=tmp_path / "outputs",
    )


def test_generate_writes_formatted_workbook_csv_and_minified_json(
    tmp_path: Path,
) -> None:
    settings = make_settings(tmp_path)
    processor = DataProcessor(
        [
            {
                "work_key": "/works/OL1W",
                "title": "The Hobbit",
                "first_publish_year": 1937,
                "languages": ["eng"],
                "author_keys": ["/authors/OL1A"],
            }
        ],
        [{"author_key": "/authors/OL1A", "name": "Tolkien"}],
    )
    paths = ReportGenerator(processor, settings).generate(on=REPORT_DATE)

    assert paths["workbook"].name == "library_report_20261008.xlsx"
    assert paths["summary"].name == "library_report_20261008.csv"
    assert paths["json"].name == "library_report_20261008.json"

    workbook = load_workbook(paths["workbook"])
    assert workbook.sheetnames == [
        "Performance Metrics",
        "Publication Years",
        "Languages",
    ]
    header = workbook["Performance Metrics"]["A1"]
    assert header.value == "metric"
    assert header.font.bold is True
    assert header.font.color.rgb[-6:] == "FFFFFF"
    assert header.fill.fgColor.rgb[-6:] == "1F4E79"
    assert workbook["Performance Metrics"].freeze_panes == "A2"
    assert workbook["Publication Years"]["A2"].value == 1937

    with paths["summary"].open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert rows[0] == {
        "summary": "publication_year",
        "label": "1937",
        "work_count": "1",
    }
    assert rows[1]["summary"] == "language"
    assert rows[1]["label"] == "eng"

    raw_json = paths["json"].read_text(encoding="utf-8")
    assert "\n" not in raw_json
    assert ": " not in raw_json
    payload = json.loads(raw_json)
    assert payload["publication_years"] == [{"publication_year": 1937, "work_count": 1}]
    assert payload["languages"][0]["language"] == "eng"


def test_empty_catalog_still_writes_all_three_files(tmp_path: Path) -> None:
    paths = ReportGenerator(DataProcessor(), make_settings(tmp_path)).generate(
        on=REPORT_DATE
    )

    assert all(path.exists() for path in paths.values())
    payload = json.loads(paths["json"].read_text(encoding="utf-8"))
    metrics = {row["metric"]: row["value"] for row in payload["metrics"]}
    assert metrics["work_count"] == 0
    assert payload["publication_years"] == []
    assert payload["languages"] == []


def test_from_database_reads_the_sqlite_catalog(tmp_path: Path) -> None:
    settings = make_settings(tmp_path)
    with DatabaseManager(settings) as database:
        database.upsert_work(
            "OL1W",
            "Dune",
            first_publish_year=1965,
            languages=["eng"],
        )
        database.upsert_author("OL2A", "Frank Herbert")
        paths = ReportGenerator.from_database(database, settings).generate(
            on=REPORT_DATE
        )

    payload = json.loads(paths["json"].read_text(encoding="utf-8"))
    metrics = {row["metric"]: row["value"] for row in payload["metrics"]}
    assert metrics["work_count"] == 1
    assert metrics["author_count"] == 1
    assert payload["publication_years"][0]["publication_year"] == 1965
